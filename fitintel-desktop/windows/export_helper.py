# -*- coding: utf-8 -*-
"""Экспорт таблиц: XLSX / PDF (DejaVu) / обёртки 1С."""
import windows.theme as theme
from paths import ASSETS, APP_DIR
from app_logging import log


def _safe_name(s: str) -> str:
    """Имя файла без запрещённых символов (для export/)."""
    import re as _re
    s = _re.sub(r'[\\/:*?"<>|]', " ", str(s or "export"))
    s = _re.sub(r"\s+", " ", s).strip().strip(".")
    return (s or "export")[:80]


def export_table_xlsx(parent, headers, rows, filename, title=""):
    from PyQt6.QtWidgets import QFileDialog, QMessageBox
    from datetime import datetime
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        QMessageBox.warning(parent, "XLSX", "Не установлен openpyxl")
        return
    export_dir = APP_DIR / "export"
    export_dir.mkdir(parents=True, exist_ok=True)
    path, _ = QFileDialog.getSaveFileName(
        parent, "Сохранить XLSX", str(export_dir / f"{_safe_name(title or filename)}.xlsx"), "Excel (*.xlsx)")
    if not path:
        return
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = _sheet_title(title or filename)
        ws.cell(1, 1, title or filename).font = Font(bold=True, size=14)
        ws.cell(2, 1, f"Дата: {datetime.now():%Y-%m-%d %H:%M}")
        try:
            user = (getattr(parent, "user_data", None) or {}).get("username", "")
            if user:
                ws.cell(3, 1, f"Пользователь: {user}")
        except Exception:
            pass
        header_row = 5
        fill = PatternFill("solid", fgColor=("1E293B" if theme.DARK else "F1F5F9"))
        font = Font(bold=True, color=("FFFFFF" if theme.DARK else "0F172A"))
        for c, h in enumerate(headers, 1):
            cell = ws.cell(header_row, c, h)
            cell.font = font
            cell.fill = fill
        for r, row in enumerate(rows, header_row + 1):
            for c, val in enumerate(row, 1):
                ws.cell(r, c, "" if val is None else str(val))
        for c in range(1, len(headers) + 1):
            max_len = 10
            for r in range(1, ws.max_row + 1):
                v = ws.cell(r, c).value
                if v:
                    max_len = max(max_len, len(str(v)))
            ws.column_dimensions[ws.cell(1, c).column_letter].width = min(max_len + 2, 60)
        wb.save(path)
        QMessageBox.information(parent, "Готово", f"Сохранено: {path}")
        log.info("XLSX сохранён: %s", path)
    except Exception as e:
        QMessageBox.warning(parent, "XLSX", f"Ошибка: {e}")
        log.error("XLSX: %s", e)


def export_table_pdf(parent, headers, rows, filename, title=""):
    from PyQt6.QtWidgets import QFileDialog, QMessageBox
    from datetime import datetime
    ttf_path = ASSETS / "DejaVuSans.ttf"
    if not ttf_path.exists():
        QMessageBox.warning(parent, "PDF",
            f"Не найден {ttf_path}\nПоложите DejaVuSans.ttf в assets/ и перезапустите клиент.")
        return
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        QMessageBox.warning(parent, "PDF", "Не установлен reportlab")
        return
    pdfmetrics.registerFont(TTFont("DejaVu", str(ttf_path)))
    export_dir = APP_DIR / "export"
    export_dir.mkdir(parents=True, exist_ok=True)
    path, _ = QFileDialog.getSaveFileName(
        parent, "Сохранить PDF", str(export_dir / f"{_safe_name(title or filename)}.pdf"), "PDF (*.pdf)")
    if not path:
        return
    try:
        doc = SimpleDocTemplate(path, pagesize=landscape(A4),
                                leftMargin=10*mm, rightMargin=10*mm,
                                topMargin=10*mm, bottomMargin=10*mm)
        story = [
            Paragraph(f"<font name='DejaVu' size='14'><b>{title or filename}</b></font>"),
            Spacer(1, 4*mm),
            Paragraph(f"<font name='DejaVu' size='9'>Дата: {datetime.now():%Y-%m-%d %H:%M}</font>"),
            Spacer(1, 6*mm),
        ]
        data = [headers] + [[str(v or "") for v in row] for row in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ]))
        story.append(table)
        doc.build(story)
        QMessageBox.information(parent, "Готово", f"Сохранено: {path}")
        log.info("PDF сохранён: %s", path)
    except Exception as e:
        QMessageBox.warning(parent, "PDF", f"Ошибка: {e}")
        log.error("PDF: %s", e)


def export_nomenclature_1c(parent, tariffs):
    from windows.accounting_tab import _1c_export_nomenclature
    _1c_export_nomenclature(parent)


def export_counterparties_1c(parent, clients):
    from windows.accounting_tab import _1c_export_counterparties
    _1c_export_counterparties(parent)


def make_export_panel(parent_widget, data_fn):
    """Панель кнопок XLSX/PDF/1С. data_fn() -> (headers, rows, title) или client_list для 1С."""
    from PyQt6.QtWidgets import QHBoxLayout, QPushButton, QMessageBox
    hb = QHBoxLayout()
    hb.addStretch()
    def _do(fmt):
        from windows.export_helper import export_table_xlsx, export_table_pdf, export_counterparties_1c
        try:
            if fmt == "1c":
                clients = data_fn("1c") if data_fn.__code__.co_argcount else []
                export_counterparties_1c(parent_widget, clients or [])
                return
            headers, rows, title = data_fn(fmt)
            if fmt == "xlsx":
                export_table_xlsx(parent_widget, headers, rows, "export", title)
            else:
                export_table_pdf(parent_widget, headers, rows, "export", title)
        except (RuntimeError, Exception) as e:
            QMessageBox.warning(parent_widget, "Экспорт", str(e)[:200])
    for fmt, label in [("xlsx", "XLSX"), ("pdf", "PDF"), ("1c", "1С")]:
        b = QPushButton(label)
        b.setStyleSheet("QPushButton { background: #f1f5f9; color: #475569; "
                        "border: 1px solid #cbd5e1; border-radius: 6px; "
                        "padding: 6px 12px; font-weight: 600; }")
        b.clicked.connect(lambda _, f=fmt: _do(f))
        hb.addWidget(b)
    return hb

def _sheet_title(s: str) -> str:
    """Имя листа XLSX: без [ ] : * ? / \\ (лимит openpyxl 31)."""
    import re as _re
    s = _re.sub(r"[\[\]:*?/\\]", " ", str(s or "export"))
    s = _re.sub(r"\s+", " ", s).strip()
    return (s or "export")[:31]


def _write_pdf_file(path, headers, rows, title):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont("DejaVu", str(ASSETS / "DejaVuSans.ttf")))
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4),
                            leftMargin=10*mm, rightMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm)
    from datetime import datetime
    story = [Paragraph(f"<font name='DejaVu' size='14'><b>{title}</b></font>"),
             Spacer(1, 4*mm),
             Paragraph(f"<font name='DejaVu' size='9'>{datetime.now():%Y-%m-%d %H:%M}</font>"),
             Spacer(1, 6*mm)]
    data = [headers] + [[str(v or "") for v in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "DejaVu"), ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
    ]))
    story.append(table)
    doc.build(story)


def save_all_formats(parent, headers, rows, title):
    """Одна кнопка — сразу XLSX+PDF+CSV+JSON в export/."""
    from PyQt6.QtWidgets import QMessageBox
    from paths import APP_DIR
    from datetime import datetime
    import json as _json, csv as _csv
    d = APP_DIR / "export"
    d.mkdir(parents=True, exist_ok=True)
    name = _safe_name(title)
    saved, errors = [], []
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        p = d / (name + ".xlsx")
        wb = Workbook(); ws = wb.active
        ws.title = _sheet_title(title)
        ws.cell(1, 1, title).font = Font(bold=True, size=14)
        ws.cell(2, 1, f"Дата: {datetime.now():%Y-%m-%d %H:%M}")
        fill = PatternFill("solid", fgColor=("1E293B" if theme.DARK else "F1F5F9"))
        fnt = Font(bold=True, color=("FFFFFF" if theme.DARK else "0F172A"))
        for c, h in enumerate(headers, 1):
            cell = ws.cell(5, c, h); cell.font = fnt; cell.fill = fill
        for r, row in enumerate(rows, 6):
            for c, val in enumerate(row, 1):
                ws.cell(r, c, "" if val is None else str(val))
        for c in range(1, len(headers) + 1):
            ml = 10
            for r in range(1, ws.max_row + 1):
                v = ws.cell(r, c).value
                if v: ml = max(ml, len(str(v)))
            ws.column_dimensions[ws.cell(1, c).column_letter].width = min(ml + 2, 60)
        wb.save(str(p)); saved.append(p)
    except Exception as e:
        errors.append("XLSX: %s" % e)
    try:
        p = d / (name + ".csv")
        with open(p, "w", newline="", encoding="utf-8-sig") as fh:
            wr = _csv.writer(fh, delimiter=";"); wr.writerow(headers); wr.writerows(rows)
        saved.append(p)
    except Exception as e:
        errors.append("CSV: %s" % e)
    try:
        p = d / (name + ".json")
        with open(p, "w", encoding="utf-8") as fh:
            _json.dump({"title": title, "exported_at": datetime.now().isoformat(timespec="seconds"),
                        "headers": headers, "rows": rows}, fh, ensure_ascii=False, indent=2)
        saved.append(p)
    except Exception as e:
        errors.append("JSON: %s" % e)
    if (ASSETS / "DejaVuSans.ttf").exists():
        try:
            p = d / (name + ".pdf")
            _write_pdf_file(p, headers, rows, title)
            saved.append(p)
        except Exception as e:
            errors.append("PDF: %s" % e)
    else:
        errors.append("PDF: нет assets/DejaVuSans.ttf")
    msg = "Сохранено файлов: %d\n%s" % (len(saved), "\n".join(str(s) for s in saved))
    if errors:
        msg += "\n\nПропущено:\n" + "\n".join(errors)
    QMessageBox.information(parent, "Сохранить (все форматы)", msg)
    return saved


def _write_xlsx_path(path, headers, rows, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from datetime import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_title(title)
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    ws.cell(2, 1, f"Дата: {datetime.now():%Y-%m-%d %H:%M}")
    fill = PatternFill("solid", fgColor=("1E293B" if theme.DARK else "F1F5F9"))
    fnt = Font(bold=True, color=("FFFFFF" if theme.DARK else "0F172A"))
    for c, h in enumerate(headers, 1):
        cell = ws.cell(5, c, h); cell.font = fnt; cell.fill = fill
    for r, row in enumerate(rows, 6):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, "" if val is None else str(val))
    for c in range(1, len(headers) + 1):
        ml = 10
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, c).value
            if v: ml = max(ml, len(str(v)))
        ws.column_dimensions[ws.cell(1, c).column_letter].width = min(ml + 2, 60)
    wb.save(str(path))


def _save_one(parent, fmt, headers, rows, title):
    """Сохранить таблицу в ОДНО выбранный формат (export/ без диалога)."""
    from PyQt6.QtWidgets import QMessageBox
    from paths import APP_DIR
    from datetime import datetime
    import json as _json, csv as _csv
    d = APP_DIR / "export"
    d.mkdir(parents=True, exist_ok=True)
    name = _safe_name(title)
    try:
        if fmt == "xlsx":
            p = d / (name + ".xlsx"); _write_xlsx_path(p, headers, rows, title)
        elif fmt == "pdf":
            if not (ASSETS / "DejaVuSans.ttf").exists():
                QMessageBox.warning(parent, "PDF", "Не найден assets/DejaVuSans.ttf"); return
            p = d / (name + ".pdf"); _write_pdf_file(p, headers, rows, title)
        elif fmt == "csv":
            p = d / (name + ".csv")
            with open(p, "w", newline="", encoding="utf-8-sig") as fh:
                wr = _csv.writer(fh, delimiter=";"); wr.writerow(headers); wr.writerows(rows)
        elif fmt == "json":
            p = d / (name + ".json")
            with open(p, "w", encoding="utf-8") as fh:
                _json.dump({"title": title, "exported_at": datetime.now().isoformat(timespec="seconds"),
                            "headers": headers, "rows": rows}, fh, ensure_ascii=False, indent=2)
        else:
            QMessageBox.warning(parent, "Сохранить", "Неизвестный формат: %s" % fmt); return
        QMessageBox.information(parent, "Сохранено", str(p))
    except Exception as e:
        QMessageBox.warning(parent, "Сохранить", str(e)[:200])


def make_save_menu_button(parent, on_format, label=" Сохранить в формате", with_1c=True):
    """Зелёная кнопка с выпадающим меню форматов. on_format(fmt) — обработчик."""
    from PyQt6.QtWidgets import QToolButton, QMenu
    btn = QToolButton()
    btn.setText(label + " ▾")
    btn.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
    btn.setStyleSheet("QToolButton { background: #10b981; color: white; border: none; "
                      "border-radius: 6px; padding: 8px 14px; font-weight: 600; }")
    menu = QMenu(btn)
    for fmt, txt in [("xlsx", "XLSX (.xlsx)"), ("pdf", "PDF (.pdf)"),
                     ("csv", "CSV (.csv)"), ("json", "JSON (.json)")]:
        act = menu.addAction(txt)
        act.triggered.connect(lambda _, f=fmt: on_format(f))
    if with_1c:
        act1c = menu.addAction("1С (XML)")
        act1c.triggered.connect(lambda _: on_format("1c"))
    btn.setMenu(menu)
    return btn


def _table_data(tbl):
    headers = [tbl.horizontalHeaderItem(i).text() for i in range(tbl.columnCount())]
    rows = [[tbl.item(r, c).text() if tbl.item(r, c) else ""
             for c in range(tbl.columnCount())] for r in range(tbl.rowCount())]
    return headers, rows


def _write_csv(parent, headers, rows, title):
    from PyQt6.QtWidgets import QFileDialog, QMessageBox
    export_dir = APP_DIR / "export"
    export_dir.mkdir(parents=True, exist_ok=True)
    import csv as _csv
    path, _ = QFileDialog.getSaveFileName(parent, "Экспорт CSV",
        str(export_dir / f"{_safe_name(title)}.csv"), "CSV (*.csv)")
    if not path:
        return
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        wr = _csv.writer(fh, delimiter=";")
        wr.writerow(headers)
        wr.writerows(rows)
    QMessageBox.information(parent, "Экспорт", f"Сохранено: {path}")


def _write_json(parent, headers, rows, title):
    from PyQt6.QtWidgets import QFileDialog, QMessageBox
    from datetime import datetime
    import json as _json
    export_dir = APP_DIR / "export"
    export_dir.mkdir(parents=True, exist_ok=True)
    path, _ = QFileDialog.getSaveFileName(parent, "Экспорт JSON",
        str(export_dir / f"{_safe_name(title)}.json"), "JSON (*.json)")
    if not path:
        return
    payload = {"title": title, "exported_at": datetime.now().isoformat(timespec="seconds"),
               "headers": headers, "rows": rows}
    with open(path, "w", encoding="utf-8") as fh:
        _json.dump(payload, fh, ensure_ascii=False, indent=2)
    QMessageBox.information(parent, "Экспорт", f"Сохранено: {path}")


def install_export_panel(widget, title=None, formats=("xlsx", "pdf", "csv", "json")):
    """Ищет самый крупный QTableWidget в widget и добавляет панель экспорта вниз."""
    from PyQt6.QtWidgets import QTableWidget, QHBoxLayout, QPushButton, QMessageBox
    if getattr(widget, "_export_panel_done", False):
        return False
    tbls = widget.findChildren(QTableWidget)
    if not tbls:
        return False
    tbl = max(tbls, key=lambda x: x.width() * x.height())
    host = tbl.parentWidget()
    lay = host.layout() if host else None
    if lay is None:
        lay = widget.layout()
    if lay is None:
        return False
    title = title or widget.__class__.__name__

    def _do(fmt):
        headers, rows = _table_data(tbl)
        try:
            if fmt == "xlsx":
                export_table_xlsx(widget, headers, rows, title, title)
            elif fmt == "pdf":
                export_table_pdf(widget, headers, rows, title, title)
            elif fmt == "csv":
                _write_csv(widget, headers, rows, title)
            elif fmt == "json":
                _write_json(widget, headers, rows, title)
            elif fmt == "1c":
                export_nomenclature_1c(widget, [])
        except (RuntimeError, Exception) as e:
            QMessageBox.warning(widget, "Экспорт", str(e)[:200])

    def _on(fmt):
        if fmt == "1c":
            export_1c_xml(widget)
            return
        if fmt == "1cml":
            export_1c_commerceml(widget)
            return
        _save_one(widget, fmt, *_table_data(tbl), title)
    btn = make_save_menu_button(widget, _on)

    hb = QHBoxLayout()
    hb.addStretch()
    hb.addWidget(btn)
    lay.addLayout(hb)
    widget._export_panel_done = True
    return True

def export_1c_xml(parent, title="export_1c"):
    """Универсальная выгрузка в 1С: XML (Export > Clients/Client, Sales/Sale), UTF-8."""
    from PyQt6.QtWidgets import QFileDialog, QMessageBox
    from paths import APP_DIR
    from xml.sax.saxutils import escape
    d = APP_DIR / "export"
    d.mkdir(parents=True, exist_ok=True)
    path, _ = QFileDialog.getSaveFileName(parent, "Выгрузка в 1С (XML)",
        str(d / "export_1c.xml"), "XML (*.xml)")
    if not path:
        return
    api = getattr(parent, "api", None)
    clients, sales = [], []
    if api is not None:
        try:
            clients = api.get_clients() or []
        except Exception:
            clients = []
        try:
            r = api._e55_get("/sales/") or {}
            sales = r.get("items", []) if isinstance(r, dict) else (r if isinstance(r, list) else [])
        except Exception:
            sales = []
    def _fio(c):
        return " ".join(x for x in (c.get("last_name"), c.get("first_name"), c.get("middle_name")) if x)
    L = ['<?xml version="1.0" encoding="UTF-8"?>', '<Export>', '  <Clients>']
    for c in clients:
        L.append('    <Client>')
        L.append('      <ExternalId>%s</ExternalId>' % escape(str(c.get("id"))))
        L.append('      <FullName>%s</FullName>' % escape(_fio(c)))
        L.append('      <Phone>%s</Phone>' % escape(str(c.get("phone") or "")))
        L.append('      <Email>%s</Email>' % escape(str(c.get("email") or "")))
        L.append('      <ContractNumber>%s</ContractNumber>' % escape(str(c.get("id"))[:8]))
        L.append('    </Client>')
    L.append('  </Clients>')
    L.append('  <Sales>')
    for s in sales[:5000]:
        amt = s.get("amount") or s.get("total") or s.get("price") or s.get("sum") or 0
        L.append('    <Sale>')
        L.append('      <ExternalId>%s</ExternalId>' % escape(str(s.get("id"))))
        L.append('      <Date>%s</Date>' % escape(str(s.get("created_at") or s.get("date") or "")[:10]))
        L.append('      <ClientExternalId>%s</ClientExternalId>' % escape(str(s.get("client_id") or "")))
        L.append('      <ProductName>%s</ProductName>' % escape(str(s.get("category") or s.get("title") or s.get("name") or s.get("description") or s.get("comment") or "Услуга")))
        L.append('      <Quantity>1</Quantity>')
        L.append('      <Price>%s</Price>' % escape(str(amt)))
        L.append('      <Total>%s</Total>' % escape(str(amt)))
        L.append('      <PaymentType>%s</PaymentType>' % escape(str(s.get("method") or s.get("payment_method") or "")))
        L.append('    </Sale>')
    L.append('  </Sales>')
    L.append('</Export>')
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))
    QMessageBox.information(parent, "1С (XML)",
        "Сохранено: %s\nКлиентов: %d, продаж: %d" % (path, len(clients), len(sales)))

def export_1c_commerceml(parent):
    """1С CommerceML с сервера: catalog/contragents/documents -> export/1c_*.xml."""
    from PyQt6.QtWidgets import QMessageBox
    from paths import APP_DIR
    api = getattr(parent, "api", None)
    if api is None:
        QMessageBox.warning(parent, "1С (CommerceML)", "Нет API-клиента")
        return
    d = APP_DIR / "export"
    d.mkdir(parents=True, exist_ok=True)
    saved, errors = [], []
    for name, path in [("catalog", "/accounting/1c/export-catalog"),
                       ("contragents", "/accounting/1c/export-contragents"),
                       ("documents", "/accounting/1c/export-documents")]:
        try:
            r = api.session.get(api._url(path), timeout=60)
            r.raise_for_status()
            p = d / ("1c_%s.xml" % name)
            p.write_bytes(r.content)
            saved.append(str(p))
        except Exception as e:
            errors.append("%s: %s" % (name, str(e)[:120]))
    msg = "Сохранено:\n" + "\n".join(saved) if saved else "Ничего не сохранено"
    if errors:
        msg += "\n\nОшибки:\n" + "\n".join(errors)
    (QMessageBox.warning if not saved else QMessageBox.information)(parent, "1С (CommerceML)", msg)